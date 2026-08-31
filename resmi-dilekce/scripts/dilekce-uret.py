#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""JSON -> resmi Türk dilekçesi (PDF + DOCX). Tek sayfa hedefler, taşarsa otomatik sıkıştırır."""
import os, sys, json, glob, subprocess

CACHE = os.path.expanduser("~/.cache/dilekce-fonts")

def font_hazirla():
    """DejaVu Serif (Türkçe tam kapsama) bul; yoksa matplotlib'ten kopyala, o da yoksa kur."""
    r, b = os.path.join(CACHE, "DejaVuSerif.ttf"), os.path.join(CACHE, "DejaVuSerif-Bold.ttf")
    if os.path.exists(r) and os.path.exists(b) and os.path.getsize(r) > 100_000:
        return r, b
    os.makedirs(CACHE, exist_ok=True)
    for attempt in (1, 2):
        try:
            import matplotlib, shutil
            d = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
            shutil.copy(os.path.join(d, "DejaVuSerif.ttf"), r)
            shutil.copy(os.path.join(d, "DejaVuSerif-Bold.ttf"), b)
            return r, b
        except Exception:
            if attempt == 2:
                sys.exit("HATA: Türkçe font bulunamadı. Çözüm: pip install matplotlib")
            subprocess.run([sys.executable, "-m", "pip", "install", "-q", "matplotlib"], check=False)
    return r, b

def uret(d, out_dir, temel_ad):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    fr, fb = font_hazirla()
    try:
        pdfmetrics.registerFont(TTFont("DJ", fr)); pdfmetrics.registerFont(TTFont("DJB", fb))
    except Exception:
        pass

    tarih   = d["tarih"]
    makam   = d["makam"] if isinstance(d["makam"], list) else [d["makam"]]
    il      = d.get("il", "")
    paras   = d["paragraflar"]
    ad      = d["ad"]
    unvan   = d.get("unvan", "")
    ekler   = d.get("ekler", [])
    iletisim= d.get("iletisim", {})   # {"T.C. Kimlik No": "...", "Telefon": "...", ...}

    os.makedirs(out_dir, exist_ok=True)
    pdf_yol = os.path.join(out_dir, temel_ad + ".pdf")

    # ---- PDF: tek sayfaya sığana kadar kademeli sıkıştır ----
    for kademe in range(4):
        fs   = [10.5, 10.5, 10.0, 9.5][kademe]
        lead = [15.5, 14.5, 14.0, 13.5][kademe]
        sp   = [1.0, 0.85, 0.7, 0.55][kademe]
        S = lambda v: Spacer(1, v * sp)
        doc = SimpleDocTemplate(pdf_yol, pagesize=A4,
                leftMargin=2.5*cm, rightMargin=2.2*cm,
                topMargin=2.2*cm, bottomMargin=1.8*cm,
                title=d.get("baslik", temel_ad), author=ad)
        st_tarih = ParagraphStyle("t", fontName="DJ",  fontSize=fs, alignment=TA_RIGHT, leading=lead)
        st_makam = ParagraphStyle("m", fontName="DJB", fontSize=fs, leading=lead+1)
        st_gov   = ParagraphStyle("g", fontName="DJ",  fontSize=fs, alignment=TA_JUSTIFY,
                                  leading=lead, firstLineIndent=1.1*cm, spaceAfter=8*sp)
        st_imza  = ParagraphStyle("s", fontName="DJ",  fontSize=fs, alignment=TA_RIGHT, leading=lead+1)
        st_k     = ParagraphStyle("k", fontName="DJ",  fontSize=fs-1, leading=lead-1)
        st_kb    = ParagraphStyle("kb",fontName="DJB", fontSize=fs-1, leading=lead-1)

        story = [Paragraph(tarih, st_tarih), S(14),
                 Paragraph("<br/>".join(makam), st_makam), S(10)]
        if il:
            story += [Paragraph(il, st_makam), S(14)]
        for p in paras:
            story.append(Paragraph(p, st_gov))
        story += [S(26), Paragraph("<b>%s</b>" % ad, st_imza)]
        if unvan:
            story.append(Paragraph(unvan, st_imza))
        story += [S(14), Paragraph("İmza: .............................", st_imza), S(20)]
        if ekler:
            story.append(Paragraph("EKİ:", st_kb))
            for i, e in enumerate(ekler, 1):
                story.append(Paragraph("%d- %s" % (i, e), st_k))
            story.append(S(14))
        if iletisim:
            story.append(Paragraph("İLETİŞİM", st_kb))
            gen = max(len(k) for k in iletisim)
            for k, v in iletisim.items():
                dolgu = "&nbsp;" * (gen - len(k)) * 2
                story.append(Paragraph("%s%s&nbsp;: %s" % (k, dolgu, v), st_k))
        doc.build(story)

        try:
            from pypdf import PdfReader
            n = len(PdfReader(pdf_yol).pages)
        except Exception:
            n = 1
        if n <= d.get("max_sayfa", 1) or kademe == 3:
            break

    # ---- DOCX ----
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc2 = docx.Document()
    sec = doc2.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.2)
    nrm = doc2.styles["Normal"]; nrm.font.name, nrm.font.size = "Times New Roman", Pt(12)
    nrm.paragraph_format.space_after = Pt(0)

    def P(t, align=None, bold=False, size=12, indent=None, after=0, ls=1.5):
        p = doc2.add_paragraph()
        if align is not None: p.alignment = align
        pf = p.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(0), ls
        if indent is not None: pf.first_line_indent = indent
        r = p.add_run(t); r.font.name, r.font.size, r.bold = "Times New Roman", Pt(size), bold

    P(tarih, WD_ALIGN_PARAGRAPH.RIGHT, after=16, ls=1)
    for line in makam: P(line, bold=True, ls=1.15)
    P("", after=10, ls=1)
    if il: P(il, bold=True, after=14, ls=1.15)
    for p in paras: P(p, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1.1), after=10)
    P("", after=26, ls=1)
    P(ad, WD_ALIGN_PARAGRAPH.RIGHT, bold=True, ls=1.15)
    if unvan: P(unvan, WD_ALIGN_PARAGRAPH.RIGHT, after=14, ls=1.15)
    P("İmza: .............................", WD_ALIGN_PARAGRAPH.RIGHT, after=22, ls=1.15)
    if ekler:
        P("EKİ:", bold=True, size=10, ls=1.15)
        for i, e in enumerate(ekler, 1): P("%d- %s" % (i, e), size=10, ls=1.15)
        P("", after=16, ls=1)
    if iletisim:
        P("İLETİŞİM", bold=True, size=10, ls=1.15)
        for k, v in iletisim.items(): P("%s\t: %s" % (k, v), size=10, ls=1.15)

    docx_yol = os.path.join(out_dir, temel_ad + ".docx")
    doc2.save(docx_yol)
    return pdf_yol, docx_yol, n

def main():
    if len(sys.argv) < 2:
        sys.exit("kullanım: dilekce-uret.py <girdi.json> [cikis_dizini]")
    d = json.load(open(sys.argv[1], encoding="utf-8"))
    out = sys.argv[2] if len(sys.argv) > 2 else d.get("cikis", ".")
    ad  = d.get("dosya_adi", "dilekce")
    pdf, dx, n = uret(d, out, ad)
    # doğrulama: metin gerçekten PDF'e girdi mi
    eksik = []
    try:
        from pypdf import PdfReader
        t = "".join(p.extract_text() for p in PdfReader(pdf).pages)
        for anahtar in d.get("dogrula", []):
            if anahtar not in t: eksik.append(anahtar)
    except Exception as e:
        print("uyarı: doğrulama yapılamadı (%s)" % e)
    print(json.dumps({"pdf": pdf, "docx": dx, "sayfa": n, "eksik": eksik},
                     ensure_ascii=False, indent=2))
    if eksik: sys.exit(4)

if __name__ == "__main__":
    main()
